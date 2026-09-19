from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("KLSCM_ABSA_Complete_Process_Documentation.docx")
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE = "E8EEF5"
LIGHT = "F4F6F9"
GRAY = "667085"
INK = "202124"
WHITE = "FFFFFF"
GOLD = "A66F00"


def font(run, size=11, bold=False, italic=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(str(text)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], PALE)
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=NAVY, size=9.2)
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    set_table_geometry(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def para(doc, text="", bold_lead=None, italic=False, color=INK, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), bold=True, color=color)
        font(p.add_run(text[len(bold_lead):]), italic=italic, color=color)
    else:
        font(p.add_run(text), italic=italic, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text))
    return p


def step(doc, title, detail):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(title + " "), bold=True, color=DARK_BLUE)
    font(p.add_run(detail))
    return p


def callout(doc, label, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    shade(t.cell(0, 0), fill)
    set_table_geometry(t, [9360])
    p = t.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(label + "  "), bold=True, color=NAVY)
    font(p.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, before, after, color in [
    ("Title", 30, 0, 8, NAVY), ("Subtitle", 14, 0, 10, GRAY),
    ("Heading 1", 16, 18, 10, BLUE), ("Heading 2", 13, 14, 7, BLUE),
    ("Heading 3", 12, 10, 5, DARK_BLUE),
]:
    s = doc.styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != "Subtitle"
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    s = doc.styles[list_name]
    s.font.name = "Calibri"
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.25

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("KLSCM MULTILINGUAL TOPIC DISCOVERY & ABSA"), size=8.5, bold=True, color=GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Project process documentation  |  Repository snapshot: 24 July 2026"), size=8.5, color=GRAY)

# Cover: editorial_cover pattern with compact_reference_guide tokens.
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("TECHNICAL PROCESS HANDBOOK"), size=10, bold=True, color=GOLD)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
font(p.add_run("KLSCM Multilingual Topic Discovery and ABSA"), size=30, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
font(p.add_run("Complete workflow, tools, data movement, controls, outputs, and operating sequence"), size=14, color=DARK_BLUE)
callout(doc, "Scope", "Instagram and online-review ingestion through multilingual preparation, relevance validation, topic discovery, aspect-based sentiment analysis, dashboards, testing, and research controls.")
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Prepared from the live repository and generated artifacts"), size=10, italic=True, color=GRAY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Snapshot date: 24 July 2026  |  Time zone: Asia/Kuala_Lumpur"), size=10, bold=True, color=GRAY)
page_break(doc)

heading(doc, "Document purpose and reading guide", 1)
para(doc, "This handbook explains what the project does, how each stage works, where code and artifacts live, when each tool is used, and which controls prevent premature or unsafe analysis. It is based on the repository source, tests, README, current processed outputs, and the Instagram cleaning report.")
callout(doc, "Important interpretation", "The repository contains both completed local preparation and later gated stages. The current processed snapshot is a 500-record relevance pilot, not a completed corpus-wide topic or ABSA result.")
heading(doc, "Contents", 2)
for item in [
    "1. System overview and current state", "2. Repository map and responsibilities",
    "3. End-to-end operating sequence", "4. Detailed stage procedures",
    "5. Tool and dependency register", "6. Data schemas and artifact lineage",
    "7. Human review and validation", "8. Dashboard and operational use",
    "9. Testing, auditability, security, and limitations", "10. Runbook and appendices",
]:
    bullet(doc, item)

heading(doc, "1. System overview and current state", 1)
para(doc, "The system prepares multilingual KLSCM text for topic discovery and aspect-based sentiment analysis (ABSA). It combines a cleaned Instagram corpus with online-review/blog records, keeps source evidence traceable, separates language-only text from sentiment-bearing text, and uses explicit gates before expensive or interpretive stages.")
heading(doc, "1.1 End-to-end flow", 2)
flow = [
    ("Raw acquisition", "Instagram JSON exports and online-review JSON/CSV"),
    ("Source cleaning", "Notebook standardizes Instagram schemas and removes blank/exact author-caption duplicates"),
    ("Unified ingestion", "Package creates stable IDs, text variants, metadata features, and duplicate audit fields"),
    ("Language preparation", "Lingua + pinned OpenLID + span analysis; optional OpenAI adjudication"),
    ("Analysis units", "Whole captions or sentence-aware review chunks, capped at 500 tokens"),
    ("Relevance gate", "Rules + OpenAI first pass + stronger-model adjudication + human review"),
    ("Validation gate", "Blind sample, repeat records, metrics, subgroup checks, target thresholds"),
    ("Topic discovery", "OpenAI embeddings + UMAP + HDBSCAN + BERTopic; taxonomy requires researcher approval"),
    ("ABSA", "Structured multilingual aspect/sentiment extraction using approved taxonomy"),
    ("Exploration", "Streamlit dashboard displays counts, language, relevance, topics, aspects, features, evidence, validation"),
]
table(doc, ["Stage", "Result"], flow, [2200, 7160])

heading(doc, "1.2 Current repository snapshot", 2)
table(doc, ["Item", "Observed state", "Meaning"], [
    ("Documents", "14,077", "14,049 Instagram + 28 blog records"),
    ("Processing status", "13,824 ready; 253 duplicate", "Duplicates remain auditable but are excluded from analysis units"),
    ("Analysis units", "14,116", "14,060 Instagram + 56 blog chunks"),
    ("Relevance pilot", "500 decisions", "318 relevant, 92 irrelevant, 90 ambiguous"),
    ("Topic inclusion", "318 Instagram documents", "Only finalized relevant decisions currently pass include_in_topics"),
    ("Pending review", "90", "Ambiguous decisions require human resolution"),
    ("Validation file", "330 rows", "300 base records + 30 repeats; 180 labels blank and malformed values are present"),
    ("Automated tests", "51 passed", "Test suite passed on 24 July 2026"),
    ("Topic/ABSA outputs", "Not present", "Expected because the relevance gate is not complete"),
], [1900, 1800, 5660])

heading(doc, "2. Repository map and responsibilities", 1)
table(doc, ["Location", "Responsibility", "Used when"], [
    ("marathon_absa/pipeline.py", "Coordinates every pipeline stage and writes run-manifest entries.", "All CLI workflows"),
    ("marathon_absa/cli.py", "Defines commands, arguments, paid-API safety flag, and console reporting.", "Operator runs the pipeline"),
    ("marathon_absa/config.py", "Central paths, pinned models, thresholds, token limits, seed, and model names.", "Imported by all stages"),
    ("marathon_absa/ingest.py", "Loads both sources and constructs canonical document records.", "prepare"),
    ("marathon_absa/text.py", "Normalization, feature extraction, text masking, hashes, stable IDs.", "ingestion and ABSA preparation"),
    ("marathon_absa/language.py", "Hybrid local language detection and mixed-language routing.", "prepare"),
    ("marathon_absa/chunking/", "Sentence-aware token chunking and unit creation.", "prepare and after language adjudication"),
    ("marathon_absa/relevance.py", "Rules, thresholds, review routing, human application, validation, topic gate.", "relevance lifecycle"),
    ("marathon_absa/openai_service.py", "Cached structured responses, embeddings, retry logic, prompts.", "paid stages"),
    ("marathon_absa/topics.py", "BERTopic fitting and initial taxonomy generation.", "topics"),
    ("marathon_absa/schemas.py", "Strict JSON schemas and initial aspect list.", "OpenAI structured stages"),
    ("marathon_absa/storage.py", "Dual Parquet/CSV output and list-column serialization.", "every persisted table"),
    ("marathon_absa/validation.py", "General validation sample and classification metrics helpers.", "validation-sample"),
    ("app.py", "Read-only analysis dashboard.", "after prepare; progressively richer later"),
    ("relevance_labeler.py", "Human relevance labeling UI with atomic CSV updates and optional translation.", "validation labeling"),
    ("tests/", "Deterministic regression and safety-gate tests.", "before changes or release"),
], [2450, 4630, 2280])

heading(doc, "3. End-to-end operating sequence", 1)
callout(doc, "Execution rule", "Run commands from the repository root in PowerShell. Paid OpenAI commands require OPENAI_API_KEY in .env and the explicit --run-api flag. Use --limit N or --pilot-size N before full-corpus processing.")
for title, detail in [
    ("Create the environment.", "Create and activate .venv, then install requirements.txt."),
    ("Download the pinned language model once.", "Run python -m marathon_absa.cli download-language-model. The exact OpenLID revision is fixed in config.py for reproducibility."),
    ("Prepare local data.", "Run python -m marathon_absa.cli prepare. This creates documents and units without paid API calls."),
    ("Review uncertain language records.", "Optionally run language-review --run-api, commonly starting with --source blog or --limit N."),
    ("Pilot relevance.", "Run relevance --run-api --pilot-size 500 to stratify by year, language, status, and text length."),
    ("Create and label the blind validation sample.", "Generate 300 records plus 30 repeats, then label relevant / ambiguous / irrelevant in the CSV or Streamlit labeler."),
    ("Validate relevance.", "Run relevance-validate only after every holdout row has a valid label; inspect precision, recall, macro F1, confusion matrix, subgroup metrics, and repeat agreement."),
    ("Classify the full Instagram corpus.", "Run relevance --run-api, resolve relevance_review_queue.csv, then run relevance-finalize."),
    ("Discover topics.", "Run topics --run-api only after validate_topic_gate confirms full coverage and zero pending reviews."),
    ("Approve the taxonomy.", "Review taxonomy.csv, enter approved_aspect values and mark approved rows before relying on the discovered taxonomy."),
    ("Run ABSA.", "Run absa --run-api; the stage uses approved aspects, or the predefined list when no approvals exist."),
    ("Explore and report.", "Run streamlit run app.py and use evidence fields plus manifests for interpretation and audit."),
]:
    step(doc, title, detail)

page_break(doc)
heading(doc, "4. Detailed stage procedures", 1)
heading(doc, "4.1 Instagram acquisition and notebook cleaning", 2)
para(doc, "The repository begins from four Instagram JSON exports collected on 10 July 2026 for #klscm2019, #klscm2023, #klscm2024, and #klscm2025. The notebook Instagram/data_cleaning.ipynb handles the historical source-specific cleaning before the reusable package pipeline begins.")
table(doc, ["Operation", "How", "Why / control"], [
    ("Schema mapping", "Maps author.id or ownerId, URL, caption, source hashtag, and timestamp to five common fields.", "Two scraper schemas are reconciled before concatenation."),
    ("Whitespace cleanup", "Trims caption edges only.", "Preserves wording, punctuation, emoji, capitalization, and internal form."),
    ("Blank removal", "Drops null/empty captions.", "Caption sentiment requires text."),
    ("Within-corpus duplicate removal", "Exact author_id + cleaned caption; first occurrence retained.", "Avoids repeated collection while allowing identical text from different authors."),
    ("Cross-corpus duplicate removal", "Same key after concatenation in 2019, 2023, 2024, 2025 order.", "A duplicate inherits the earliest source corpus in that order."),
    ("Timestamp normalization", "Parses UTC-aware timestamps and asserts success.", "Creates a consistent temporal reference."),
    ("Export verification", "Writes UTF-8-SIG CSV without index, reads it back, checks columns and row count.", "Protects multilingual text and spreadsheet compatibility."),
], [2200, 3650, 3510])
para(doc, "Cleaning reduced 14,591 raw Instagram records to 14,049. The source hashtag is a retrieval-corpus label, not a strict publication-year label; timestamps extend beyond the named event editions.")

heading(doc, "4.2 Unified ingestion and text representations", 2)
para(doc, "Pipeline.prepare calls load_documents. Instagram comes from Instagram/instagram_cleanded.csv; blog reviews come from Online Review Blog/raw-data.json. Each source record is transformed into one canonical document.")
table(doc, ["Representation", "Transformation", "Downstream use"], [
    ("original_text", "Exact input string.", "Evidence, display, traceability"),
    ("normalized_text", "ftfy repair + Unicode NFKC + collapsed whitespace.", "Deduplication and chunking"),
    ("linguistic_text", "Removes URLs, @mentions, hashtags, KLSCM metadata tokens, and emoji.", "Language detection without metadata bias"),
    ("semantic_text", "Removes URLs/mentions, converts emoji to aliases, converts hashtags to hashtag_* tokens.", "Retains sentiment-bearing social features"),
    ("hashtags / emojis / aliases", "Extracted as list-valued audit features.", "Dashboard and later explanatory analysis"),
    ("text_hash", "SHA-256 of case-folded word-like canonical text.", "Exact normalized duplicate detection"),
    ("document_id", "First 20 hex characters of a stable SHA-256 over source identifiers and ordinal.", "Joins and repeatable lineage"),
], [2000, 4160, 3200])
callout(doc, "Preservation policy", "Empty and duplicate records are not deleted. processing_status and duplicate_of preserve the audit trail; only ready records become analysis units.")

heading(doc, "4.3 Hybrid language detection", 2)
para(doc, "Language detection is deliberately separated from semantic analysis. The service first uses Lingua locally. When the pinned OpenLID-v3 model exists, OpenLID confirms uncertain, multilingual, Malay/Indonesian-sensitive, or multi-sentence cases and analyzes sentences plus overlapping 16-word windows.")
table(doc, ["Decision point", "Configured rule"], [
    ("Minimum content", "At least 8 linguistic letters after metadata removal"),
    ("Local confidence", "0.80 minimum"),
    ("Top-two margin", "0.15 minimum"),
    ("Mixed-language component", "At least 20% coverage, 10 characters, and 0.35 component confidence"),
    ("Sliding windows", "16 words with 5-word overlap"),
    ("Malay shorthand", "Known forms are normalized; two or more hits provide a Malay hint"),
    ("Non-core languages", "Need stronger confidence, length, and Lingua corroboration"),
], [3000, 6360])
para(doc, "High-confidence Lingua results can complete locally. OpenLID produces primary and secondary language, confidence, margin, span coverage, mixed-language status, and the detection method. Missing OpenLID falls back to Lingua but marks the result review_required and pending.")
table(doc, ["Status", "Meaning / next action"], [
    ("no_text", "Normalized source is empty; no analysis."),
    ("insufficient_text", "Fewer than 8 linguistic characters; excluded from language percentages."),
    ("ok", "Accepted single-language local decision."),
    ("mixed", "Two or more substantial language spans; optional OpenAI adjudication queued."),
    ("low_confidence", "Confidence or margin failed; queued."),
    ("review_required", "Detector disagreement, Malay/Indonesian ambiguity, or fallback; queued."),
    ("undetermined", "No credible supported language; queued."),
], [2300, 7060])

heading(doc, "4.4 Optional OpenAI language adjudication", 2)
para(doc, "language-review reads pending, ready documents, optionally filters by source and limit, and sends local findings plus original text to a strict JSON schema. The result replaces final language fields with language_method=openai_adjudicated, records confidence and reasons, writes language_reviews, rewrites documents, and rebuilds units so downstream language metadata stays synchronized.")

heading(doc, "4.5 Sentence-aware chunking", 2)
para(doc, "The SentenceChunker uses tiktoken for the configured chat model, falling back to o200k_base and finally a byte-based estimate. Instagram captions at or below 500 tokens remain whole. Blog records and longer captions are split on multilingual sentence punctuation and packed without exceeding 500 tokens. Oversized single sentences are token-sliced.")
table(doc, ["Field", "Purpose"], [
    ("unit_id", "Stable document + chunk ordinal key"),
    ("sentence_start / sentence_end", "Original sentence-position lineage"),
    ("token_count", "Cost and boundary audit"),
    ("text", "Exact analysis unit sent to embedding or ABSA"),
    ("document/source/year/language/url", "Context copied for joins and evidence"),
], [3000, 6360])

heading(doc, "4.6 Relevance classification and routing", 2)
para(doc, "Relevance is applied only to ready Instagram documents. Blog reviews bypass this Instagram relevance filter and are included later when ready. The event policy covers registration, event-specific preparation, logistics, participant information, participation, experience, results, achievement, support, and evaluation.")
table(doc, ["Layer", "When used", "Outcome"], [
    ("Deterministic rule", "Insufficient or hashtag-only linguistic content.", "Ambiguous, pending human review; no API call"),
    ("Initial OpenAI model", "All other candidate captions.", "Structured relevant / ambiguous / irrelevant decision"),
    ("Stronger adjudication model", "Initial label is ambiguous, relevant confidence < 0.80, or irrelevant confidence < 0.90.", "Independent second pass"),
    ("Routing thresholds", "After adjudication.", "Relevant if >=0.75; irrelevant if >=0.90; otherwise ambiguous/pending"),
    ("Human review", "All pending rows in relevance_review_queue.csv.", "Final relevant or irrelevant overrides model output"),
], [2450, 3300, 3610])
para(doc, "Each record stores evidence, English gloss, observed language, code-switching note, initial and adjudication model/prompt/cache/token metadata, review status, adjudication method, and include_in_topics. Existing human decisions are preserved across reruns.")
callout(doc, "Single inclusion authority", "include_in_topics is the only Instagram inclusion gate for topics and ABSA. Ambiguous, unreviewed, empty, duplicate, and uncovered records cannot silently enter.")

heading(doc, "4.7 Human relevance validation", 2)
para(doc, "relevance-review-sample builds a stratified blind file across event year, language, language status, and length band. The default contains 300 original records plus 30 repeated records. The separate key retains model labels and a calibration/holdout split.")
table(doc, ["Measure", "Required target / role"], [
    ("Relevant recall", ">= 0.90; prioritizes retaining genuine event content"),
    ("Relevant precision", ">= 0.85; limits contamination"),
    ("Macro F1", ">= 0.80 across relevant, ambiguous, irrelevant"),
    ("Confusion matrix", "Shows error direction"),
    ("Subgroup metrics", "Year, primary language, language status, length band"),
    ("Intra-reviewer Cohen's kappa", "Agreement between originals and hidden repeats"),
], [3100, 6260])
callout(doc, "Current blocker", "The live validation file is not ready for validation: 180 labels are blank and malformed labels exist. Correct these entries before running relevance-validate.")

heading(doc, "4.8 Topic discovery", 2)
para(doc, "topics first calls validate_topic_gate. The gate requires relevance coverage for every eligible Instagram document and zero pending reviews. It then selects include_in_topics Instagram units plus all ready blog units.")
table(doc, ["Tool / setting", "How it is used"], [
    ("OpenAI embeddings", "text-embedding-3-large by default; batches of 256 units"),
    ("UMAP", "Cosine metric, 5 components, min_dist 0, fixed random seed; neighbors capped at 15"),
    ("HDBSCAN", "Euclidean clustering, EOM selection, prediction data enabled"),
    ("BERTopic", "Multilingual topic representation; needs at least 10 units"),
    ("Dynamic cluster size", "max(5, min(30, number of units // 50))"),
], [3100, 6260])
para(doc, "Outputs are topic_assignments, topic_info, a pickled BERTopic model, and taxonomy.csv. Outlier topic -1 is excluded from the initial taxonomy. A researcher supplies approved_aspect, description, approval flag, and notes.")

heading(doc, "4.9 Aspect-based sentiment analysis", 2)
para(doc, "absa reuses the same topic gate and selected units. It loads approved taxonomy aspects; when no usable approvals exist, it falls back to the predefined 13-aspect list, always retaining other_emerging. Each unit is sent unchanged for multilingual structured analysis; translation is not performed before classification.")
table(doc, ["Output field", "Meaning"], [
    ("aspect / emerging_aspect", "Approved category or named new category"),
    ("target", "Entity or feature evaluated"),
    ("sentiment", "positive, neutral, negative, or mixed"),
    ("confidence", "0-1 model confidence"),
    ("evidence", "Exact source span"),
    ("english_gloss", "Short post-analysis gloss"),
    ("aspect_expression", "explicit or implicit"),
    ("contributing_hashtags / emojis", "Social features that contribute meaning"),
    ("model_observed_language / code_switching_note", "Multilingual audit context"),
], [3200, 6160])
para(doc, "Mention IDs are stable hashes of unit, mention ordinal, aspect, and evidence. Results merge incrementally into aspect_mentions; API metadata is written to absa_processing_log.")

heading(doc, "5. Tool and dependency register", 1)
table(doc, ["Tool", "What it does here", "When / where used"], [
    ("pandas", "Tabular ingestion, cleaning, joining, sampling, exports.", "Notebook and nearly every package stage"),
    ("pyarrow", "Parquet engine.", "storage.write_table/read_table"),
    ("ftfy + unicodedata", "Repairs mojibake and applies NFKC.", "text.repair_and_normalize"),
    ("emoji", "Extracts emoji and aliases; removes or demojizes by representation.", "text.py and dashboard"),
    ("Lingua", "Fast local confidence-based language detection.", "prepare"),
    ("fastText + OpenLID-v3", "Pinned broad language model and span predictions.", "prepare after one-time download"),
    ("Hugging Face Hub", "Downloads exact OpenLID file/revision.", "download-language-model"),
    ("tiktoken", "Token counts and chunk boundaries.", "prepare chunking"),
    ("OpenAI Responses API", "Strict JSON language review, relevance, ABSA.", "paid commands with --run-api"),
    ("OpenAI Embeddings API", "Semantic vectors.", "topics --run-api"),
    ("BERTopic / UMAP / HDBSCAN", "Dimensionality reduction, density clustering, topic summaries.", "topics"),
    ("scikit-learn", "Validation metrics and kappa.", "relevance validation"),
    ("Streamlit", "Dashboard and human labeler.", "app.py and relevance_labeler.py"),
    ("Plotly", "Interactive charts.", "app.py"),
    ("pytest", "Deterministic regression tests.", "development and release checks"),
    ("python-dotenv", "Loads .env from repository root.", "config import"),
    ("Pydantic", "Declared dependency for typed validation ecosystem.", "Available; schemas currently use JSON-schema dictionaries"),
], [2250, 3900, 3210])

heading(doc, "6. Data schemas and artifact lineage", 1)
table(doc, ["Artifact", "Producer", "Consumer / purpose"], [
    ("Instagram/instagram_cleanded.csv", "Instagram notebook", "Unified ingestion"),
    ("Online Review Blog/raw-data.json", "Source collection", "Unified ingestion"),
    ("data/processed/documents.parquet + .csv", "prepare / language-review", "Dashboard, relevance, units, validation"),
    ("data/processed/units.parquet + .csv", "prepare / language-review", "Topics and ABSA"),
    ("data/cache/<stage>/<hash>.json", "CachedOpenAI", "Repeatable paid calls and cost avoidance"),
    ("relevance.parquet + .csv", "relevance / relevance-finalize", "Gate, dashboard, validation, topics, ABSA"),
    ("relevance_review_queue.*", "relevance", "Human resolution"),
    ("relevance_validation_sample.csv", "relevance-review-sample", "Blind human labels"),
    ("relevance_validation_key.*", "relevance-review-sample", "Held-back model key"),
    ("relevance_* metrics files", "relevance-validate", "Quality decision"),
    ("topic_assignments.* / topic_info.*", "topics", "Dashboard and taxonomy work"),
    ("taxonomy.csv", "topics then researcher", "ABSA aspect allow-list"),
    ("bertopic_model.pkl", "topics", "Reproducibility / later inspection"),
    ("aspect_mentions.*", "absa", "Evidence and aspect-sentiment dashboard"),
    ("absa_processing_log.*", "absa", "Model, prompt, cache, latency, token audit"),
    ("run_manifest.jsonl", "pipeline stages", "Append-only run history"),
], [2850, 2500, 4010])

heading(doc, "7. Human review workflows", 1)
heading(doc, "7.1 Relevance labeling application", 2)
para(doc, "Run streamlit run relevance_labeler.py after creating relevance_validation_sample.csv. The app validates the required schema, navigates to pending records, records relevant / ambiguous / irrelevant, displays context, and optionally translates languages outside the reviewer-ready set using the cached OpenAI service.")
para(doc, "Each save creates one timestamped session backup, writes a temporary CSV in the same directory, flushes and fsyncs it, then atomically replaces the target with os.replace. Invalid labels are rejected without changing the file. Translation preserves names, handles, hashtags, emoji, URLs, line breaks, and KLSCM terms and is never used as the classification source.")
heading(doc, "7.2 Review rules", 2)
for item in [
    "Judge semantic connection in the caption itself; do not infer unavailable image content.",
    "A KLSCM hashtag alone is not evidence of event relevance.",
    "Use relevant for the full event journey, not only race-day opinions.",
    "Use irrelevant for unrelated events, generic running, pure promotion, lifestyle/spam, or no connection.",
    "Use ambiguous when evidence is genuinely insufficient, image-dependent, weak, or conflicting.",
    "Complete review_notes when a difficult judgment needs an audit explanation.",
    "Resolve all malformed labels before validation; allowed values are exact lowercase labels.",
]:
    bullet(doc, item)

heading(doc, "8. Dashboard and operational use", 1)
para(doc, "app.py is a read-only Streamlit dashboard. It loads optional Parquet artifacts and degrades gracefully when later stages are absent. Filters cover source, event year, primary language, mixed-language status, and detection method.")
table(doc, ["Tab", "What the operator sees", "Available after"], [
    ("Overview", "Documents, ready count, chunks, duplicates, source/year charts.", "prepare"),
    ("Language & relevance", "Language distribution, quality queues, mixed records, relevance funnel and reasons.", "prepare; relevance adds more"),
    ("Topics", "Topic counts and topic information.", "topics"),
    ("Aspect sentiment", "Aspect-by-sentiment counts.", "absa"),
    ("Hashtags & emoji", "Top features by source.", "prepare"),
    ("Evidence", "Aspect-level records and exact evidence.", "absa"),
    ("Validation", "General validation sample and run manifest.", "validation-sample / any manifested run"),
], [1800, 5000, 2560])
callout(doc, "Interpretation safeguard", "Language charts exclude empty, duplicate, no-text, and insufficient-text records. Hashtag-year strata must not be described as strict event-year publication samples without an additional date filter.")

heading(doc, "9. Testing, auditability, security, and limitations", 1)
heading(doc, "9.1 Automated tests", 2)
para(doc, "The current suite contains deterministic tests for language masking, mixed-language behavior, pipeline safety gates, relevance routing, human review preservation, topic inclusion, labeler navigation, atomic saving, translation contracts, and malformed input. The full command .\\.venv\\Scripts\\python.exe -m pytest -q passed 51 tests on this snapshot.")
heading(doc, "9.2 Audit and reproducibility controls", 2)
for item in [
    "Pinned OpenLID repository revision in config.py.",
    "Fixed random seed 42 for sampling and UMAP.",
    "Stable SHA-256-derived document, unit, review, and mention IDs.",
    "Prompt versions and model names stored with API outputs.",
    "Content-addressed JSON cache for structured OpenAI requests.",
    "Append-only UTC run manifest for major stages.",
    "Dual Parquet and UTF-8-SIG CSV outputs for machine use and inspection.",
    "Raw records and exclusion reasons retained rather than deleted.",
]:
    bullet(doc, item)
heading(doc, "9.3 Security and cost controls", 2)
for item in [
    "OPENAI_API_KEY belongs only in .env and must never be logged or committed.",
    "Paid CLI stages refuse to run unless --run-api is explicitly supplied.",
    "--limit and --pilot-size support bounded trial runs.",
    "Cached responses prevent paying twice for identical stage/model/prompt/text inputs.",
    "API calls retry up to five times with exponential backoff.",
    "data/cache, data/processed, models, .env, and virtual environments are local generated artifacts and should remain uncommitted.",
]:
    bullet(doc, item)
heading(doc, "9.4 Known methodological and implementation limitations", 2)
for item in [
    "Instagram hashtag sampling is not a complete population of KLSCM discourse.",
    "Collection in July 2026 is affected by deletion, privacy, availability, platform ranking, and scraper behavior.",
    "Exact normalized-text duplicates do not identify paraphrases or copied text across materially different wording.",
    "Unified package deduplication is text-hash based across sources, while the notebook uses author + caption; these are distinct layers and should be reported accurately.",
    "Language decisions can be difficult for short, romanized, slang-heavy, or closely related Malay/Indonesian text.",
    "OpenAI classifications and embeddings are model-dependent; prompt/model versions must accompany reported results.",
    "BERTopic clusters are exploratory structures, not self-validating research constructs; taxonomy approval is required.",
    "The dashboard reports available artifacts and does not itself enforce that validation targets passed.",
    "Current code creates a general validation sample separately from the relevance-specific blind validation workflow; the two should not be confused.",
]:
    bullet(doc, item)

heading(doc, "10. Operator runbook", 1)
table(doc, ["Goal", "PowerShell command", "Precondition / expected output"], [
    ("Set up", "python -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npip install -r requirements.txt", "Creates local environment"),
    ("Test", "pytest", "All deterministic tests pass; no paid APIs"),
    ("Download OpenLID", "python -m marathon_absa.cli download-language-model", "One-time ~1.2 GB pinned model"),
    ("Prepare", "python -m marathon_absa.cli prepare", "documents.* and units.*"),
    ("Language pilot", "python -m marathon_absa.cli language-review --run-api --source blog --limit 20", "Pending language records; paid"),
    ("Relevance pilot", "python -m marathon_absa.cli relevance --run-api --pilot-size 500", "relevance.* and review queue; paid"),
    ("Blind sample", "python -m marathon_absa.cli relevance-review-sample --size 300 --repeats 30", "Validation sample and secret key"),
    ("Human labeler", "streamlit run relevance_labeler.py", "Edits validation CSV atomically"),
    ("Validate", "python -m marathon_absa.cli relevance-validate", "All holdout labels valid and complete"),
    ("Full relevance", "python -m marathon_absa.cli relevance --run-api", "Full eligible Instagram coverage; paid"),
    ("Finalize", "python -m marathon_absa.cli relevance-finalize", "Human queue resolved"),
    ("Topics", "python -m marathon_absa.cli topics --run-api", "Gate complete; topic artifacts; paid"),
    ("ABSA", "python -m marathon_absa.cli absa --run-api", "Gate complete and taxonomy reviewed; paid"),
    ("Dashboard", "streamlit run app.py", "Reads whatever processed artifacts exist"),
], [1600, 4100, 3660])

heading(doc, "10.1 Recommended immediate next actions for this snapshot", 2)
for item in [
    "Clean the malformed human_relevance entries and complete the 180 blank validation labels.",
    "Run relevance-validate and examine both overall targets and subgroup recall.",
    "If targets fail, revise the relevance policy/prompt/thresholds and repeat a controlled pilot rather than scaling.",
    "After targets pass, run full-corpus relevance, complete the review queue, and finalize.",
    "Run topics, review and approve taxonomy.csv, then run ABSA.",
    "Capture dashboard screenshots and validation metrics when preparing a research report or pull request.",
]:
    bullet(doc, item)

heading(doc, "Appendix A. Configuration reference", 1)
table(doc, ["Setting", "Default", "Operational meaning"], [
    ("OPENAI_CHAT_MODEL", "gpt-5.6-luna", "Language review, initial relevance, ABSA"),
    ("OPENAI_STRONGER_MODEL", "gpt-5.6-terra", "Borderline relevance adjudication"),
    ("OPENAI_EMBEDDING_MODEL", "text-embedding-3-large", "Topic vectors"),
    ("OpenLID repo", "HPLT/OpenLID-v3", "Local multilingual detection"),
    ("OpenLID revision", "6b9560483e17e42f48d86cebf22b4b58dffeaa70", "Exact reproducible model revision"),
    ("Minimum language letters", "8", "Below this becomes insufficient_text"),
    ("Language confidence / margin", "0.80 / 0.15", "Local acceptance controls"),
    ("Mixed coverage / characters", "0.20 / 10", "Substantial span definition"),
    ("Language window / overlap", "16 / 5 words", "Span analysis"),
    ("Chunk target / maximum", "400 / 500 tokens", "Target reserved; maximum enforced"),
    ("Random seed", "42", "Sampling and UMAP reproducibility"),
], [2850, 3000, 3510])

heading(doc, "Appendix B. Initial ABSA aspect vocabulary", 1)
for aspect in [
    "route_scenery", "weather_climate", "event_organization", "registration_communication",
    "expo_race_kit", "transport_accessibility", "crowd_atmosphere",
    "aid_stations_refreshments", "facilities_amenities", "safety_medical",
    "cost_value", "personal_race_experience", "other_emerging",
]:
    bullet(doc, aspect)

heading(doc, "Appendix C. Source basis", 1)
para(doc, "This document was derived from README.md; AGENTS.md; app.py; relevance_labeler.py; requirements.txt; all marathon_absa modules; tests; Instagram/data_cleaning_report.md; and the current data/processed artifact inventory. It describes observed code behavior as of 24 July 2026 and does not claim that absent topic or ABSA artifacts have been produced.")

doc.core_properties.title = "KLSCM Multilingual Topic Discovery and ABSA - Complete Process Documentation"
doc.core_properties.subject = "Technical process handbook"
doc.core_properties.author = "Project documentation generated from repository evidence"
doc.core_properties.keywords = "KLSCM, multilingual, BERTopic, ABSA, relevance, language detection"
doc.save(OUT)
print(OUT.resolve())
